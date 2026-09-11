"""Exportacion a Excel y PDF.

En Excel el tiempo se escribe como valor de tiempo (minutos / 1440) con formato
[h]:mm, para que las sumas funcionen dentro de la hoja. En PDF se muestra H:MM.
"""

import io

from django.http import HttpResponse
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.tiempo import excel_tiempo

FORMATO_TIEMPO = "[h]:mm"
RELLENO_TITULO = PatternFill("solid", fgColor="1F3864")
FUENTE_TITULO = Font(color="FFFFFF", bold=True)


def hoja_nueva(titulo: str):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    return libro, hoja


def escribir_encabezado(hoja, columnas: list[str], fila: int = 1) -> None:
    for i, texto in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila, column=i, value=texto)
        celda.fill = RELLENO_TITULO
        celda.font = FUENTE_TITULO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.freeze_panes = hoja.cell(row=fila + 1, column=1)


def escribir_tiempo(hoja, fila: int, columna: int, minutos: int) -> None:
    celda = hoja.cell(row=fila, column=columna, value=excel_tiempo(minutos))
    celda.number_format = FORMATO_TIEMPO


def ajustar_anchos(hoja, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def respuesta_excel(libro, nombre: str) -> HttpResponse:
    flujo = io.BytesIO()
    libro.save(flujo)
    respuesta = HttpResponse(
        flujo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return respuesta


def respuesta_pdf(plantilla: str, contexto: dict, nombre: str) -> HttpResponse:
    """Genera el PDF a partir de la misma plantilla HTML que se ve en pantalla."""
    from xhtml2pdf import pisa

    html = render_to_string(plantilla, {**contexto, "para_pdf": True})
    flujo = io.BytesIO()
    error = pisa.CreatePDF(html, dest=flujo, encoding="utf-8")
    if error.err:
        return HttpResponse("No se pudo generar el PDF.", status=500)
    respuesta = HttpResponse(flujo.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}.pdf"'
    return respuesta
