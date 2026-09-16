"""El reporte de asistencia: una fila por dia con las marcas en columnas.

Tiene que servir desde el primer dia, sin empleados ni horarios creados: solo
con lo que manda el reloj. Los casos salen de marcas reales de setiembre.
"""

import io
from datetime import date, time, timedelta

import pytest
from openpyxl import load_workbook

from apps.core.tiempo import datetime_local
from apps.marcas.models import MarcaReloj
from apps.marcas.servicio import crear_marca_manual, ingestar
from apps.reportes import dias
from apps.motor import calculo
from tests.conftest import marca_json, utc_ms_de

MARTES = date(2026, 9, 15)


def marcar(person_id, nombre, fecha, *horas):
    ingestar([
        marca_json(person_id, utc_ms_de(fecha, h), person_name=nombre) for h in horas
    ])


def marca(fecha, h, origen=calculo.RELOJ):
    return calculo.Marca(hora=datetime_local(fecha, h), origen=origen, ref_id=0)


# --------------------------------------------------------------------------
# Armado de un dia
# --------------------------------------------------------------------------


def test_jornada_partida_de_benjamin():
    """Informe de SmartPSS del 15/09: cuatro marcas, dos tramos."""
    d = dias.armar_dia(MARTES, [
        marca(MARTES, time(16, 30)), marca(MARTES, time(5, 45)),
        marca(MARTES, time(20, 26)), marca(MARTES, time(14, 5)),
    ])
    assert [(e.hora.strftime("%H:%M"), s.hora.strftime("%H:%M")) for e, s in d.pares] == [
        ("05:45", "14:05"), ("16:30", "20:26"),
    ]
    assert d.minutos == 500 + 236
    assert d.completo
    assert d.observacion == ""


def test_una_sola_marca_queda_incompleta_y_no_suma_horas():
    d = dias.armar_dia(MARTES, [marca(MARTES, time(5, 49))])
    assert not d.completo
    assert d.minutos == 0
    assert "Falta la salida de las 05:49" in d.observacion


def test_las_repetidas_no_rompen_el_par():
    """Roque Aguilar, 23/08: 11:49 20:49 20:49. Sin descartar, parecia incompleto."""
    d = dias.armar_dia(MARTES, [
        marca(MARTES, time(11, 49)), marca(MARTES, time(20, 49)), marca(MARTES, time(20, 49)),
    ])
    assert d.completo
    assert d.minutos == 540
    assert d.repetidas == 1
    assert "repetida" in d.observacion


def test_hacen_falta_mas_columnas_si_alguien_marca_mas():
    persona = dias.Persona(clave="1", person_id="1", nombre="X", dias=[
        dias.armar_dia(MARTES, [marca(MARTES, time(h, 0)) for h in (6, 8, 10, 12, 14, 16)])
    ])
    assert dias.columnas_de_marcas([persona]) == 3
    assert dias.columnas_de_marcas([]) == 2


# --------------------------------------------------------------------------
# Pantalla y Excel
# --------------------------------------------------------------------------


@pytest.fixture
def cliente(client, usuario):
    client.force_login(usuario)
    return client


@pytest.fixture
def con_marcas(db):
    marcar("14", "BENJAMIN QUIROS MORA", MARTES, time(5, 45), time(14, 5), time(16, 30), time(20, 26))
    marcar("16", "BRAYAN JORGE SOLANO GUIDO", MARTES, time(5, 49))
    marcar("9", "ADRIANA MARIA CASTRO AZOFEIFA", MARTES, time(5, 50), time(14, 1))


@pytest.mark.django_db
def test_sale_sin_empleados_ni_horarios(cliente, con_marcas):
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15"})
    assert r.status_code == 200
    html = r.content.decode()
    for texto in ("BENJAMIN QUIROS MORA", "05:45", "20:26", "12:16",
                  "BRAYAN JORGE SOLANO GUIDO", "Falta la salida de las 05:49"):
        assert texto in html
    assert r.context["total_incompletos"] == 1


@pytest.mark.django_db
def test_filtrar_solo_incompletos(cliente, con_marcas):
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15", "incompletos": "1"})
    assert [p.nombre for p in r.context["personas"]] == ["BRAYAN JORGE SOLANO GUIDO"]


@pytest.mark.django_db
def test_buscar_por_nombre(cliente, con_marcas):
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15", "buscar": "adriana"})
    assert [p.person_id for p in r.context["personas"]] == ["9"]


@pytest.mark.django_db
def test_una_marca_manual_completa_el_dia_y_se_senala(cliente, maria, usuario):
    marcar("1024", "Maria", MARTES, time(8, 0))
    crear_marca_manual(maria, datetime_local(MARTES, time(17, 0)), "olvido", "Lo vio el guarda.", usuario)

    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15"})
    persona = r.context["personas"][0]
    assert persona.dias[0].completo
    assert persona.minutos == 540
    assert "manual" in persona.dias[0].observacion


@pytest.mark.django_db
def test_una_marca_anulada_no_sale(cliente, con_marcas, usuario):
    from apps.marcas.servicio import anular_marca

    anular_marca(MarcaReloj.objects.get(person_id="16"), usuario, "Marco por error")
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15"})
    assert "16" not in [p.person_id for p in r.context["personas"]]


@pytest.mark.django_db
def test_el_excel_trae_las_mismas_columnas_y_horas_sumables(cliente, con_marcas):
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15", "formato": "excel"})
    assert r.status_code == 200
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Detalle", "Resumen"]

    hoja = libro["Detalle"]
    encabezado = [c.value for c in hoja[4]]
    assert encabezado[:9] == ["PersonID", "Codigo", "Nombre", "Fecha", "Dia", "E1", "S1", "E2", "S2"]
    col_horas = encabezado.index("Horas") + 1

    filas = {hoja.cell(row=i, column=3).value: i for i in range(5, hoja.max_row + 1)}
    benjamin = filas["BENJAMIN QUIROS MORA"]
    assert [hoja.cell(row=benjamin, column=c).value for c in range(6, 10)] == ["05:45", "14:05", "16:30", "20:26"]
    horas = hoja.cell(row=benjamin, column=col_horas)
    assert horas.number_format == "[h]:mm"
    # Excel la guarda como duracion, no como texto: por eso se puede sumar.
    assert horas.value == timedelta(minutes=736)


# --------------------------------------------------------------------------
# Corregir un dia incompleto desde el reporte
# --------------------------------------------------------------------------


@pytest.mark.django_db
def test_el_dia_incompleto_trae_boton_para_corregir(cliente, con_marcas):
    r = cliente.get("/reportes/", {"desde": "2026-09-15", "hasta": "2026-09-15"})
    html = r.content.decode()
    assert "/marcas/dia/16/2026-09-15/?volver=" in html
    assert "Corregir" in html


@pytest.mark.django_db
def test_sin_horario_la_pantalla_del_dia_muestra_entrada_y_falta_de_salida(cliente, con_marcas):
    """Antes todas las marcas salian como 'no cuenta' y no se veia que corregir."""
    r = cliente.get("/marcas/dia/16/2026-09-15/")
    assert r.status_code == 200
    [marca] = r.context["marcas_reloj"]
    assert marca.papel == "entrada"
    assert marca.sin_pareja is True
    assert r.context["sin_horario"].completo is False
    assert "falta la salida" in r.content.decode()


@pytest.mark.django_db
def test_agregar_la_salida_completa_el_dia_y_vuelve_al_reporte(cliente, con_marcas):
    volver = "/reportes/?desde=2026-09-15&hasta=2026-09-15"
    r = cliente.post("/marcas/dia/16/2026-09-15/manual/", {
        "hora": "14:00", "motivo": "olvido", "detalle": "Confirmado con el supervisor.", "volver": volver,
    })
    assert r.status_code == 302
    assert "volver=" in r["Location"]

    reporte = cliente.get(volver)
    brayan = next(p for p in reporte.context["personas"] if p.person_id == "16")
    assert brayan.dias[0].completo
    assert brayan.minutos == 8 * 60 + 11
    assert "manual" in brayan.dias[0].observacion


@pytest.mark.django_db
@pytest.mark.parametrize("malo", ["https://otro.com/", "//otro.com/", "javascript:alert(1)"])
def test_el_enlace_de_volver_no_lleva_fuera_del_sistema(cliente, con_marcas, malo):
    r = cliente.get("/marcas/dia/16/2026-09-15/", {"volver": malo})
    assert r.context["volver"] == ""